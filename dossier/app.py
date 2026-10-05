import hmac
import tkinter as tk
import webbrowser
from pathlib import Path
from tkinter import messagebox
from urllib.parse import urlsplit
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException


APP_DIR = Path(__file__).resolve().parent
ACCESS_FILE = APP_DIR / "utilisateurs_autorises.xlsx"
FORM_URL_FILE = APP_DIR / "lien_formulaire.txt"
ACCESS_HEADERS = ("Identifiant", "Code d'accès")


def create_access_template(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.create_sheet("Acces")
    workbook.remove(workbook.worksheets[0])
    sheet.append(ACCESS_HEADERS)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 28
    for row in range(2, 1002):
        sheet.cell(row=row, column=1).number_format = "@"
        sheet.cell(row=row, column=2).number_format = "@"
    workbook.save(path)
    workbook.close()


def has_authorized_access(path: Path, identifier: str, access_code: str) -> bool:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if "Acces" not in workbook.sheetnames:
            raise ValueError("L'onglet « Acces » est introuvable dans le fichier Excel.")

        sheet = workbook["Acces"]
        headers = tuple(sheet.cell(row=1, column=column).value for column in (1, 2))
        if headers != ACCESS_HEADERS:
            raise ValueError(
                "Les deux premières cellules de la ligne 1 doivent être : "
                "Identifiant et Code d'accès."
            )

        normalized_identifier = identifier.strip().casefold()
        normalized_code = access_code.strip()
        for stored_identifier, stored_code in sheet.iter_rows(
            min_row=2, max_col=2, values_only=True
        ):
            if stored_identifier is None or stored_code is None:
                continue
            if (
                str(stored_identifier).strip().casefold() == normalized_identifier
                and hmac.compare_digest(
                    str(stored_code).strip().encode("utf-8"),
                    normalized_code.encode("utf-8"),
                )
            ):
                return True
        return False
    finally:
        workbook.close()


def read_form_url(path: Path) -> str:
    url = path.read_text(encoding="utf-8").strip()
    parsed_url = urlsplit(url)
    if parsed_url.scheme != "https" or not parsed_url.hostname:
        raise ValueError(
            "Le fichier lien_formulaire.txt doit contenir une URL complète "
            "commençant par https://."
        )
    return url


class LoginMockup(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Bienvenue")
        self.geometry("620x700")
        self.minsize(460, 620)
        self.configure(bg="#f1f3f6")

        if not ACCESS_FILE.exists():
            create_access_template(ACCESS_FILE)

        card = tk.Frame(self, bg="white", padx=54, pady=42)
        card.pack(fill="both", expand=True, padx=18, pady=18)

        tk.Label(
            card,
            text="Bienvenue",
            bg="white",
            fg="#202b38",
            font=("Segoe UI", 27, "bold"),
            anchor="w",
        ).pack(fill="x")

        tk.Label(
            card,
            text="Connectez-vous à votre espace",
            bg="white",
            fg="#748398",
            font=("Segoe UI", 13),
            anchor="w",
        ).pack(fill="x", pady=(12, 24))

        tk.Label(
            card,
            text=(
                "Utilisez l'identifiant et le code d'accès dédiés inscrits dans "
                "utilisateurs_autorises.xlsx. N'utilisez pas votre mot de passe "
                "de compte."
            ),
            bg="#fff4e5",
            fg="#704b16",
            font=("Segoe UI", 9),
            justify="left",
            wraplength=440,
            padx=12,
            pady=10,
        ).pack(fill="x", pady=(0, 24))

        tk.Label(
            card, text="Identifiant", bg="white", fg="#202b38", font=("Segoe UI", 11)
        ).pack(fill="x", pady=(0, 8))
        self.identifier = tk.Entry(
            card,
            font=("Segoe UI", 14),
            relief="solid",
            bd=1,
            highlightthickness=1,
            highlightbackground="#d5d9de",
            highlightcolor="#7892ad",
        )
        self.identifier.pack(fill="x", ipady=9, pady=(0, 22))

        tk.Label(
            card,
            text="Code d'accès",
            bg="white",
            fg="#202b38",
            font=("Segoe UI", 11),
        ).pack(fill="x", pady=(0, 8))
        self.access_code = tk.Entry(
            card,
            font=("Segoe UI", 14),
            show="•",
            relief="solid",
            bd=1,
            highlightthickness=1,
            highlightbackground="#d5d9de",
            highlightcolor="#7892ad",
        )
        self.access_code.pack(fill="x", ipady=9)

        self.show_code = tk.BooleanVar(value=False)
        tk.Checkbutton(
            card,
            text="Afficher le code d'accès",
            variable=self.show_code,
            command=self.toggle_code,
            bg="white",
            activebackground="white",
            fg="#202b38",
            font=("Segoe UI", 10),
            anchor="w",
        ).pack(fill="x", pady=(10, 28))

        tk.Button(
            card,
            text="CONNEXION",
            command=self.open_form,
            bg="white",
            fg="#e45d63",
            activebackground="#fff5f5",
            activeforeground="#cf454b",
            font=("Segoe UI", 13),
            relief="solid",
            bd=1,
            cursor="hand2",
        ).pack(fill="x", ipady=4)

    def toggle_code(self):
        self.access_code.configure(show="" if self.show_code.get() else "•")

    def open_form(self):
        identifier = self.identifier.get().strip()
        access_code = self.access_code.get()
        self.identifier.delete(0, tk.END)
        self.access_code.delete(0, tk.END)

        if not identifier or not access_code.strip():
            messagebox.showwarning(
                "Champs requis",
                "Saisissez votre identifiant et votre code d'accès.",
                parent=self,
            )
            return

        try:
            authorized = has_authorized_access(
                ACCESS_FILE, identifier, access_code
            )
        except (OSError, BadZipFile, InvalidFileException, ValueError) as error:
            messagebox.showerror("Fichier Excel invalide", str(error), parent=self)
            return

        if not authorized:
            messagebox.showerror(
                "Accès refusé",
                "Cet identifiant et ce code d'accès ne figurent pas dans le fichier "
                "des utilisateurs autorisés.",
                parent=self,
            )
            return

        try:
            form_url = read_form_url(FORM_URL_FILE)
        except (OSError, ValueError) as error:
            messagebox.showerror("Lien du formulaire invalide", str(error), parent=self)
            return

        try:
            opened = webbrowser.open_new_tab(form_url)
        except webbrowser.Error as error:
            messagebox.showerror(
                "Ouverture impossible",
                f"Le formulaire n'a pas pu être ouvert : {error}",
                parent=self,
            )
            return

        if not opened:
            messagebox.showerror(
                "Ouverture impossible",
                "Le navigateur n'a pas pu ouvrir le formulaire.",
                parent=self,
            )


if __name__ == "__main__":
    LoginMockup().mainloop()
